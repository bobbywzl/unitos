"""The sheets benchmark's helpers in Python (scripts/parse-bench/sheets.mts),
never the code under test:

  csv <file> <out.json> <encoding> <delimiter>
      The reference for a .csv/.tsv: Python's csv module reads the text in
      the encoding the corpus records (the file's true encoding, checked by
      hand), RFC 4180 quoting. One sheet; trailing empty cells and rows
      trimmed; control characters other than tab and newline dropped (they
      are not visible, and the parse drops them too, SPEC.md §5).
  derive <src> <out> <src-encoding> <src-delimiter> <encoding> <delimiter> <eol> <decimal-comma 0|1>
      A real file saved the way another spreadsheet program saves it: the
      rows read with Python's csv module and written again with the new
      delimiter, line end, and encoding; with decimal-comma, a plain decimal
      number's point becomes a comma (a German or French Excel's CSV).
  strict <in.xlsx> <out.xlsx>
      A Strict OOXML workbook (purl.oclc.org namespaces) written with the
      transitional namespaces, so POI can read it (POI bug 57699). Writes
      nothing when the workbook is not strict.
  synth <id> <out.xlsx>
      The synthetic workbooks: sizes and layouts no public file in the
      corpus has (openpyxl, fixed seed).
"""
import csv
import io
import json
import random
import re
import sys
import zipfile

CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def read_rows(path, encoding, delimiter):
    with open(path, "rb") as f:
        data = f.read()
    text = data.decode(encoding)
    if text.startswith("﻿"):
        text = text[1:]
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', doublequote=True, strict=False))


def csv_ref(path, out, encoding, delimiter):
    rows = []
    for row in read_rows(path, encoding, delimiter):
        cells = [CONTROL.sub("", v) for v in row]
        while cells and cells[-1] == "":
            cells.pop()
        rows.append([{"t": v, "k": "s"} if v != "" else None for v in cells])
    while rows and not rows[-1]:
        rows.pop()
    sheet = {"name": None, "hidden": False, "rows": rows, "merges": [], "frozenRows": 0, "frozenCols": 0}
    if len(rows) > 10000:
        sheet["cutRows"] = len(rows)
        sheet["rows"] = rows[:10000]
    if any(len(r) > 256 for r in sheet["rows"]):
        sheet["rows"] = [r[:256] for r in sheet["rows"]]
    with open(out, "w", encoding="utf-8") as f:
        json.dump({"sheets": [sheet]}, f, ensure_ascii=False)


DECIMAL = re.compile(r"^-?\d+\.\d+$")


def derive(src, out, src_encoding, src_delimiter, encoding, delimiter, eol, decimal_comma):
    rows = read_rows(src, src_encoding, src_delimiter)
    if decimal_comma:
        rows = [[v.replace(".", ",") if DECIMAL.match(v) else v for v in row] for row in rows]
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=delimiter, quotechar='"', lineterminator=eol, quoting=csv.QUOTE_MINIMAL)
    writer.writerows(rows)
    with open(out, "wb") as f:
        f.write(buf.getvalue().encode(encoding))


STRICT_NS = re.compile(rb"http://purl\.oclc\.org/ooxml/(\w+)/(\w+)")


def strict(src, out):
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        workbook = [n for n in names if n.endswith("workbook.xml")]
        if not any(b"purl.oclc.org/ooxml" in z.read(n) for n in workbook):
            return
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as w:
            for n in names:
                data = z.read(n)
                if n.endswith(".xml") or n.endswith(".rels"):
                    data = STRICT_NS.sub(rb"http://schemas.openxmlformats.org/\1/2006/\2", data)
                    data = data.replace(b' conformance="strict"', b"")
                w.writestr(n, data)


def synth(id, out):
    import datetime
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    rnd = random.Random(1)
    wb = openpyxl.Workbook()
    ws = wb.active
    words = ["north", "south", "east", "west", "river", "field", "stone", "harbor", "market", "mill"]
    if id == "synth-long":
        ws.title = "Records"
        ws.append(["Id", "Name", "Region", "Date", "Amount", "Share", "Count", "Code", "Ratio", "Flag", "Note", "Total"])
        for r in range(2, 15002):
            ws.append([
                r - 1,
                f"{rnd.choice(words).title()} {rnd.choice(words)}",
                rnd.choice(words),
                datetime.date(2020, 1, 1) + datetime.timedelta(days=rnd.randint(0, 1500)),
                round(rnd.uniform(-5000, 50000), 2),
                rnd.random(),
                rnd.randint(0, 999),
                f"C-{rnd.randint(1000, 9999)}",
                rnd.uniform(0, 3),
                rnd.random() < 0.5,
                rnd.choice(words) if rnd.random() < 0.3 else None,
                f"=E{r}*G{r}",
            ])
            row = ws[r]
            row[3].number_format = "yyyy-mm-dd"
            row[4].number_format = "#,##0.00"
            row[5].number_format = "0.0%"
            row[8].number_format = "0.000"
        ws.freeze_panes = "B2"
    elif id == "synth-wide":
        ws.title = "Wide"
        for r in range(1, 41):
            ws.append([f"r{r}c{c}" if (r + c) % 7 else rnd.randint(1, 10 ** 6) for c in range(1, 301)])
        ws.freeze_panes = "A2"
    elif id == "synth-frozen-merged":
        ws.title = "Plan"
        ws["A1"] = "Quarterly plan"
        ws.merge_cells("A1:F1")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.append(["(hidden note)", "", "", "", "", ""])
        ws.row_dimensions[2].hidden = True
        ws.append(["Team", "Lead", "Q1", "Q2", "Q3", "Q4"])
        ws.column_dimensions["B"].hidden = True
        for i, team in enumerate(["Design", "Build", "Ship", "Support", "Sales", "Finance"]):
            ws.append([team, f"Lead {i}", rnd.uniform(0, 1), rnd.uniform(0, 1), rnd.randint(1000, 90000), rnd.randint(-500, 500)])
            r = ws.max_row
            ws.cell(r, 3).number_format = "0%"
            ws.cell(r, 4).number_format = "0.00%"
            ws.cell(r, 5).number_format = '"$"#,##0'
            ws.cell(r, 6).number_format = "#,##0;[Red](#,##0)"
        ws.append(["Notes", "", "Line one\nLine two", "", "", ""])
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=ws.max_row, start_column=3, end_row=ws.max_row, end_column=6)
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="FFEEAA")
        ws.append(["Total", "", "=AVERAGE(C4:C9)", "=AVERAGE(D4:D9)", "=SUM(E4:E9)", "=SUM(F4:F9)"])
        ws.merge_cells("A12:A13")
        ws["A12"] = "Merged down"
        ws["B13"] = "under the hidden column"
        ws["C13"] = 0.5
        ws["C13"].number_format = "# ?/?"
        ws["D13"] = 44197.75
        ws["D13"].number_format = "d mmm yyyy h:mm AM/PM"
        ws.freeze_panes = "C4"
        side = wb.create_sheet("Hidden data")
        side.sheet_state = "hidden"
        side["A1"] = "never shown"
        last = wb.create_sheet("After hidden")
        last["A1"] = "shown after a hidden sheet"
    elif id == "synth-merged-hidden-words":
        # A timetable whose merged headers keep words in the cells they
        # cover, as files written by libraries keep them (openpyxl drops
        # them on merge, so the parts are written as they are).
        return merged_hidden_words(out)
    elif id == "synth-time-rounding":
        # A shift log whose times sit a fraction of a second under a
        # minute, an hour, or a day, as times summed or read off a clock
        # are stored.
        ws.title = "Shifts"
        ws.append(["Entry", "Time", "Format"])
        for label, value, fmt in [
            ("Start", 0.3645833, "hh:mm:ss"),
            ("Clock in", 0.36458, "hh:mm"),
            ("End of day", 0.999999, "hh:mm"),
            ("Last second", 0.999999, "hh:mm:ss"),
            ("Two days", 1.9999999, "[h]:mm:ss"),
            ("Stamp", 44197.9999999, "yyyy-mm-dd hh:mm:ss"),
            ("Break", 0.0208333, "h:mm:ss AM/PM"),
            ("Noon", 0.5, "h:mm AM/PM"),
            ("Shift", 1 / 3, "h:mm:ss"),
        ]:
            ws.append([label, value, fmt])
            ws.cell(ws.max_row, 2).number_format = fmt
    else:
        raise SystemExit(f"unknown synth id {id}")
    wb.save(out)


def merged_hidden_words(out):
    def cell(ref, text):
        text = text.replace("&", "&amp;").replace("<", "&lt;")
        return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'

    rows = [
        ["A1:Time", "B1:Monday", "C1:left over", "D1:Tuesday", "E1:draft"],
        ["A2:9:00", "B2:Maths", "C2:Room 4", "D2:History", "E2:Room 2"],
        ["A3:10:00", "B3:Break", "C3:hidden", "D3:Science", "E3:Lab"],
        ["A4:11:00", "B4:stale", "C4:Room 4", "D4:Art", "E4:Room 9"],
        ["A5:Notes", "B5:Bring the forms", "C5:covered", "D5:covered too", "E5:"],
    ]
    sheet_rows = []
    for i, row in enumerate(rows, 1):
        cells = "".join(cell(*c.split(":", 1)) for c in row if c.split(":", 1)[1])
        sheet_rows.append(f'<row r="{i}">{cells}</row>')
    # Monday and Tuesday span their room columns; the break spans two
    # rows; the notes span the row.
    merges = ["B1:C1", "D1:E1", "B3:B4", "B5:E5"]
    merge_cells = "".join('<mergeCell ref="%s"/>' % m for m in merges)
    data = "".join(sheet_rows)
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook {ns}><sheets><sheet name="Timetable" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet {ns}><sheetData>{data}</sheetData><mergeCells count="{len(merges)}">{merge_cells}</mergeCells></worksheet>',
    }
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def main(argv):
    cmd = argv[0]
    if cmd == "csv":
        csv_ref(argv[1], argv[2], argv[3], argv[4])
    elif cmd == "derive":
        derive(argv[1], argv[2], argv[3], argv[4], argv[5], argv[6], argv[7], argv[8] == "1")
    elif cmd == "strict":
        strict(argv[1], argv[2])
    elif cmd == "synth":
        synth(argv[1], argv[2])
    else:
        raise SystemExit(f"unknown command {cmd}")


if __name__ == "__main__":
    main(sys.argv[1:])
